"""Check reader-facing appendix claims against the public gold source, without a solver."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def audit(root: Path = ROOT) -> dict:
    """Read frozen source cells and metadata; do not prepare or solve an instance."""
    manuscript = root / "manuscript"
    supplement = manuscript / "supplement"
    guide = json.loads(
        (root / "docs/evidence/appendix_input_reproduction_20261028.json").read_text()
    )
    inventory = json.loads((supplement / "appendix_d_inventory.json").read_text())
    gold = root / guide["public_gold"]["repository_path"]
    digest = hashlib.sha256(gold.read_bytes()).hexdigest()
    if digest != guide["public_gold"]["sha256"]:
        raise ValueError("Public gold bytes differ from the identified source")
    formulas = openpyxl.load_workbook(gold, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(gold, read_only=True, data_only=True)
    try:
        if len(formulas.sheetnames) != guide["public_gold"]["worksheet_count"]:
            raise ValueError("Gold worksheet count differs")
        for cell in inventory["reference_cells"]:
            sheet, address = cell["sheet"], cell["cell"]
            if (
                formulas[sheet][address].value != cell["original"]
                or cached[sheet][address].value != cell["cached"]
            ):
                raise ValueError(f"Original/cached source cell differs: {sheet}!{address}")
        rows = list(cached["Warehouses"].iter_rows(values_only=True))
        headers = rows[0]
        records = [
            dict(zip(headers, row, strict=True))
            for row in rows[1:] if row[0] is not None
        ]
        existing = [r for r in records if str(r["Status"]).casefold().startswith("exist")]
        candidates = [r for r in records if str(r["Status"]).casefold().startswith("candidat")]
        expected = guide["public_gold"]["structural_counts"]
        if (len(records), len(existing), len(candidates)) != (
            expected["warehouses"], expected["existing"], expected["candidates"]
        ):
            raise ValueError("Canonical population roles differ")
        freights = {
            cached["Frete"].cell(row, 1).value: cached["Frete"].cell(row, 2).value
            for row in range(2, 29)
        }
        if freights != {
            r["state"]: r["value"] for r in guide["freight_brl_per_tonne_km"]
        }:
            raise ValueError("Reader guide freight values differ from actual cells")
        if guide["executed_scalar_contract"]["penalty_rates"] != {
            "unmet": 90000, "static": 45000, "reception": 45000
        }:
            raise ValueError("Executed penalty vector differs")
        registry = json.loads(
            (root / "data/manifests/warehouse_population_source_v020.json").read_text()
        )
        if registry["source_sha256"] != guide["registry"]["sha256"]:
            raise ValueError("Registry identity differs")
        for case in guide["processed_workbooks"]:
            matches = [
                w for w in inventory["workbooks"]
                if w.get("population") == case["population"]
            ]
            if len(matches) != 1 or any(
                case[k] != matches[0][k]
                for k in ("sha256", "existing", "candidate", "ordered_ids_sha256")
            ):
                raise ValueError("Processed workbook or membership identity differs")
        links = []
        files = [supplement / name for name in (
            "index.qmd", "appendix_c.qmd", "appendix_d.qmd"
        )]
        prefix = "https://github.com/unb-lamfo-or-ai-hpc/model-agrologistic/blob/"
        for source in files:
            for target in re.findall(r"\]\((https://[^)]+)\)", source.read_text()):
                if not target.startswith(prefix):
                    continue
                ref, _, relative = target.removeprefix(prefix).partition("/")
                path = root / relative.split("#", 1)[0]
                if not re.fullmatch(r"[0-9a-f]{40}", ref) or not path.is_file():
                    raise ValueError(f"Nonexistent or unpinned appendix source link: {target}")
                links.append(target)
        sample_fields = (
            "CDA", "Status", "Cap. Estática (t)", "Cap. Recepção (t)",
            "Cap. Expedição (t)", "Cap. Estática Máxima (t)"
        )
        return {
            "scope": "Public source and appendix-link audit; no model, solver or private-input access",
            "gold_sha256": digest,
            "gold_worksheets": len(formulas.sheetnames),
            "original_and_cached_cells_checked": len(inventory["reference_cells"]),
            "warehouses": len(records),
            "existing": len(existing),
            "candidates": len(candidates),
            "freight_states_checked": len(freights),
            "pinned_github_file_links_checked": len(links),
            "source_examples_not_executed_decisions": [
                {key: record[key] for key in sample_fields}
                for record in (existing[0], candidates[0])
            ],
            "external_registry_bytes_reread": False,
            "processed_workbook_bytes_reread": False,
            "large_optimization_repeated": False,
        }
    finally:
        formulas.close()
        cached.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = json.dumps(audit(), indent=2, ensure_ascii=False) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(result, encoding="utf-8")
    print(result, end="")
