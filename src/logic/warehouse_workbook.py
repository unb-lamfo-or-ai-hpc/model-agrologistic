"""Materialize nested policy workbooks from an audited warehouse population."""

from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from src.logic.warehouse_population import (
    MAX_STATIC_CAPACITY,
    OPENING_COST,
    RECEPTION_CAPACITY,
    SHIPPING_CAPACITY,
    STATIC_CAPACITY,
    STATUS,
    WAREHOUSE_ID,
    _population_role,
)

CANDIDATE_FIXED_COST = "Custo_Fixo_Abertura_Modelo ($)"
CANDIDATE_VARIABLE_COST = "Custo_Capacidade_Candidata_Modelo ($/t)"
EXPANSION_ALLOWED = "Permite_Expansao"
MAX_EXPANSION = "Cap_Expansao_Maxima_Modelo (t)"
EXPANSION_VARIABLE_COST = "Custo_Expansao_Modelo ($/t)"
BULKIFICATION_ALLOWED = "Permite_Granelizacao"
BULKIFICATION_VARIABLE_COST = "Custo_Granelizacao_Modelo ($/t)"
PARAMETER_SOURCE = "Fonte_Parametro"
OBSERVATION = "Observacao"
WORKBOOK_CHANGE = "Alteracao_Excel_Modelo"


@dataclass(frozen=True, slots=True)
class PopulationWorkbookArtifacts:
    """Paths created by one population-workbook materialization."""

    output_root: Path
    manifest_json: Path
    workbooks: tuple[Path, ...]


def build_population_workbooks(
    source_workbook: Path,
    anchor_workbook: Path,
    population_order_csv: Path,
    output_root: Path,
    *,
    target_populations: tuple[int, ...] | None = None,
    source_sheet_name: str = "Sheet1",
    anchor_sheet_name: str = "Warehouses",
    overwrite: bool = False,
) -> PopulationWorkbookArtifacts:
    """Build population-only workbooks while preserving candidate activation."""

    source_workbook = Path(source_workbook)
    anchor_workbook = Path(anchor_workbook)
    population_order_csv = Path(population_order_csv)
    output_root = Path(output_root)

    source = pd.read_excel(source_workbook, sheet_name=source_sheet_name)
    anchor = pd.read_excel(anchor_workbook, sheet_name=anchor_sheet_name)
    order = pd.read_csv(population_order_csv)
    targets = _resolve_targets(order, target_populations)
    _validate_materialization_inputs(source, anchor, order, targets)

    source_by_id = source.set_index(WAREHOUSE_ID, drop=False)
    anchor_by_id = anchor.set_index(WAREHOUSE_ID, drop=False)
    existing_count = int((order["population_role"] == "existing").sum())
    output_root.mkdir(parents=True, exist_ok=True)

    records: list[dict[str, Any]] = []
    workbooks: list[Path] = []
    for target in targets:
        selected = order.loc[
            (order["population_role"] == "existing")
            | (
                order["population_eligible"].map(_as_bool)
                & order["candidate_rank"].le(target - existing_count)
            )
        ].copy()
        if len(selected) != target:
            raise ValueError(
                f"Population {target} selected {len(selected)} warehouses."
            )

        rows = _materialize_rows(selected, source_by_id, anchor_by_id, anchor.columns)
        output_path = output_root / f"warehouses_{target}" / "model_input.xlsx"
        if output_path.exists() and not overwrite:
            raise FileExistsError(
                f"Population workbook already exists: {output_path}."
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(anchor_workbook, output_path)
        _replace_warehouse_sheet(output_path, anchor_sheet_name, anchor.columns, rows)
        validation = _validate_materialized_workbook(
            output_path,
            anchor_sheet_name,
            selected,
            source_by_id,
            anchor_by_id,
        )
        records.append(
            {
                "target_population": target,
                "workbook": str(output_path.relative_to(output_root)),
                "workbook_sha256": _sha256(output_path),
                **validation,
            }
        )
        workbooks.append(output_path)

    manifest_path = output_root / "population_workbook_manifest.json"
    manifest = {
        "schema_version": 1,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "source_workbook": source_workbook.name,
        "source_workbook_sha256": _sha256(source_workbook),
        "anchor_workbook": anchor_workbook.name,
        "anchor_workbook_sha256": _sha256(anchor_workbook),
        "population_order_sha256": _sha256(population_order_csv),
        "distance_status": "pending_osrm_materialization",
        "candidate_translation": {
            "base_static_capacity": 0.0,
            "base_reception_capacity": 0.0,
            "base_shipping_capacity": 0.0,
            "maximum_static_capacity": "observed source static capacity",
            "investment_cost": "derived from the canonical Custo_Invest table",
        },
        "populations": records,
    }
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return PopulationWorkbookArtifacts(
        output_root=output_root,
        manifest_json=manifest_path,
        workbooks=tuple(workbooks),
    )


def _resolve_targets(
    order: pd.DataFrame,
    requested: tuple[int, ...] | None,
) -> tuple[int, ...]:
    available = tuple(
        int(value)
        for value in sorted(order["first_population_total"].dropna().unique())
    )
    targets = (
        available
        if requested is None
        else tuple(int(value) for value in requested)
    )
    if not targets:
        raise ValueError("At least one target population is required.")
    unknown = sorted(set(targets) - set(available))
    if unknown:
        raise ValueError(f"Targets are absent from the population order: {unknown}.")
    if tuple(sorted(set(targets))) != targets:
        raise ValueError("Target populations must be strictly increasing.")
    return targets


def _validate_materialization_inputs(
    source: pd.DataFrame,
    anchor: pd.DataFrame,
    order: pd.DataFrame,
    targets: tuple[int, ...],
) -> None:
    required_order_columns = {
        WAREHOUSE_ID,
        "population_role",
        "population_eligible",
        "candidate_rank",
        "first_population_total",
    }
    missing_order = sorted(required_order_columns - set(order.columns))
    if missing_order:
        raise ValueError(f"Population order is missing columns: {missing_order}.")
    required_warehouse_columns = {
        WAREHOUSE_ID,
        STATUS,
        STATIC_CAPACITY,
        RECEPTION_CAPACITY,
        SHIPPING_CAPACITY,
        MAX_STATIC_CAPACITY,
        OPENING_COST,
    }
    for name, frame in (("source", source), ("anchor", anchor)):
        missing = sorted(required_warehouse_columns - set(frame.columns))
        if missing:
            raise ValueError(f"The {name} workbook is missing columns: {missing}.")
        if frame[WAREHOUSE_ID].astype(str).duplicated().any():
            raise ValueError(f"The {name} workbook has duplicate CDA identifiers.")
    if targets[-1] > len(order.loc[order["population_eligible"].map(_as_bool)]):
        raise ValueError("The requested population exceeds eligible facilities.")


def _materialize_rows(
    selected: pd.DataFrame,
    source_by_id: pd.DataFrame,
    anchor_by_id: pd.DataFrame,
    columns: pd.Index,
) -> list[dict[str, Any]]:
    rows = []
    for warehouse in selected[WAREHOUSE_ID].astype(str):
        if warehouse in anchor_by_id.index:
            rows.append(anchor_by_id.loc[warehouse].to_dict())
            continue
        if warehouse not in source_by_id.index:
            raise ValueError(f"Warehouse {warehouse!r} is absent from the registry.")
        source_row = source_by_id.loc[warehouse]
        translated = {
            column: source_row[column] if column in source_row.index else None
            for column in columns
        }
        if _population_role(translated[STATUS]) != "candidate":
            raise ValueError(
                f"Noncanonical warehouse {warehouse!r} must be a candidate."
            )
        observed_static_capacity = float(source_row[STATIC_CAPACITY])
        if observed_static_capacity <= 0.0:
            raise ValueError(
                f"Candidate {warehouse!r} has nonpositive source capacity."
            )
        translated[STATIC_CAPACITY] = 0.0
        translated[RECEPTION_CAPACITY] = 0.0
        translated[SHIPPING_CAPACITY] = 0.0
        translated[MAX_STATIC_CAPACITY] = observed_static_capacity
        translated[OPENING_COST] = None
        translated[CANDIDATE_FIXED_COST] = 0.0
        translated[CANDIDATE_VARIABLE_COST] = None
        translated[EXPANSION_ALLOWED] = "NO"
        translated[MAX_EXPANSION] = None
        translated[EXPANSION_VARIABLE_COST] = None
        translated[BULKIFICATION_ALLOWED] = "NO"
        translated[BULKIFICATION_VARIABLE_COST] = None
        translated[PARAMETER_SOURCE] = (
            "Observed capacity from the audited registry; candidate investment "
            "cost from Custo_Invest."
        )
        translated[OBSERVATION] = (
            "Candidate base capacity is zero until the opening decision is active."
        )
        translated[WORKBOOK_CHANGE] = (
            "Added by the v0.2.0 nested warehouse population protocol."
        )
        rows.append(translated)
    return rows


def _replace_warehouse_sheet(
    workbook_path: Path,
    sheet_name: str,
    columns: pd.Index,
    rows: list[dict[str, Any]],
) -> None:
    workbook = openpyxl.load_workbook(workbook_path)
    worksheet = workbook[sheet_name]
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    for row in rows:
        worksheet.append([_excel_value(row.get(column)) for column in columns])
    workbook.save(workbook_path)


def _validate_materialized_workbook(
    workbook_path: Path,
    sheet_name: str,
    selected: pd.DataFrame,
    source_by_id: pd.DataFrame,
    anchor_by_id: pd.DataFrame,
) -> dict[str, Any]:
    result = pd.read_excel(workbook_path, sheet_name=sheet_name)
    selected_ids = selected[WAREHOUSE_ID].astype(str).tolist()
    result_ids = result[WAREHOUSE_ID].astype(str).tolist()
    if result_ids != selected_ids:
        raise ValueError("Materialized warehouse order differs from the audit order.")

    anchor_ids = set(anchor_by_id.index) & set(selected_ids)
    new_candidate_ids = set(selected_ids) - set(anchor_by_id.index)
    result_by_id = result.set_index(WAREHOUSE_ID, drop=False)
    for warehouse in anchor_ids:
        for column in anchor_by_id.columns:
            if not _values_equal(
                result_by_id.loc[warehouse, column],
                anchor_by_id.loc[warehouse, column],
            ):
                raise ValueError(
                    f"Canonical warehouse {warehouse!r} changed column {column!r}."
                )
    candidates = result.loc[result[STATUS].map(_population_role) == "candidate"]
    zero_base = candidates[
        [STATIC_CAPACITY, RECEPTION_CAPACITY, SHIPPING_CAPACITY]
    ].fillna(0.0).eq(0.0).all().all()
    positive_maximum = candidates[MAX_STATIC_CAPACITY].fillna(0.0).gt(0.0).all()
    if not zero_base or not positive_maximum:
        raise ValueError("Candidate activation translation did not validate.")

    for warehouse in new_candidate_ids:
        observed = float(source_by_id.loc[warehouse, STATIC_CAPACITY])
        materialized = float(
            result.loc[result[WAREHOUSE_ID] == warehouse, MAX_STATIC_CAPACITY].iloc[0]
        )
        if materialized != observed:
            raise ValueError(
                f"Candidate {warehouse!r} did not preserve observed capacity."
            )
    return {
        "warehouse_count": int(len(result)),
        "existing_warehouse_count": int(
            (result[STATUS].map(_population_role) == "existing").sum()
        ),
        "candidate_warehouse_count": int(len(candidates)),
        "canonical_anchor_row_count": int(len(anchor_ids)),
        "additional_candidate_count": int(len(new_candidate_ids)),
        "candidate_activation_contract_valid": True,
    }


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "true", "yes"}


def _excel_value(value: Any) -> Any:
    return None if pd.isna(value) else value


def _values_equal(left: Any, right: Any) -> bool:
    if pd.isna(left) and pd.isna(right):
        return True
    if isinstance(left, float) or isinstance(right, float):
        try:
            return abs(float(left) - float(right)) <= 1e-9
        except (TypeError, ValueError):
            return False
    return left == right


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

