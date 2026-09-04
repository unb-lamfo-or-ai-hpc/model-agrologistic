from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.logic.excel_loader import ExcelLoaderConfig, load_model_data_from_excel
from src.logic.model_config import ModelConfig
from src.logic.model_validation import validate_model_data


def build_tiny_golden_excel(path: Path) -> None:
    oferta = pd.DataFrame(
        [
            {
                "Produto": "Soja",
                "Cidade": "Rio Verde - GO",
                "Latitude": -17.7923,
                "Longitude": -50.9192,
                "Data": "2026-01",
                "Peso (ton)": 100.0,
            },
            {
                "Produto": "Soja",
                "Cidade": "Jataí - GO",
                "Latitude": -17.8784,
                "Longitude": -51.7204,
                "Data": "2026-01",
                "Peso (ton)": 50.0,
            },
            {
                "Produto": "Soja",
                "Cidade": "Rio Verde - GO",
                "Latitude": -17.7923,
                "Longitude": -50.9192,
                "Data": "2026-02",
                "Peso (ton)": 120.0,
            },
        ]
    )

    demanda = pd.DataFrame(
        [
            {
                "Produto": "Soja",
                "Cidade": "Goiânia - GO",
                "Latitude": -16.6869,
                "Longitude": -49.2648,
                "Data": "2026-01",
                "Peso (ton)": 80.0,
                "Tipo_Demanda": "DOMESTICA",
                "Regra_Limite": "FIXO",
                "Peso_Modelo (ton)": 80.0,
                "Fonte_Parametro": "Informado",
                "Observacao": "Demanda doméstica fixa",
            },
            {
                "Produto": "Soja",
                "Cidade": "Santos - SP",
                "Latitude": -23.9535,
                "Longitude": -46.3350,
                "Data": "2026-01",
                "Peso (ton)": None,
                "Tipo_Demanda": "EXPORTACAO",
                "Regra_Limite": "AUTO_OFERTA_TOTAL_PRODUTO_PERIODO",
                "Peso_Modelo (ton)": None,
                "Fonte_Parametro": "Derivado_da_Oferta",
                "Observacao": "Mercado externo não restritivo",
            },
            {
                "Produto": "Soja",
                "Cidade": "Santos - SP",
                "Latitude": -23.9535,
                "Longitude": -46.3350,
                "Data": "2026-02",
                "Peso (ton)": None,
                "Tipo_Demanda": "EXPORTACAO",
                "Regra_Limite": "AUTO_OFERTA_TOTAL_PRODUTO_PERIODO",
                "Peso_Modelo (ton)": None,
                "Fonte_Parametro": "Derivado_da_Oferta",
                "Observacao": "Mercado externo não restritivo",
            },
        ]
    )

    warehouses = pd.DataFrame(
        [
            {
                "CDA": "W1",
                "Status": "Existente",
                "Município": "Rio Verde",
                "UF": "GO",
                "Latitude": -17.7923,
                "Longitude": -50.9192,
                "Armazenador": "COMPANHIA NACIONAL DE ABASTECIMENTO",
                "Tipo": "Convencional",
                "Cap. Estática (t)": 200.0,
                "Cap. Recepção (t)": 200.0,
                "Cap. Expedição (t)": 200.0,
                "Cap. Estática Máxima (t)": 0.0,
                "Custo de Abertura ($)": 0.0,
                "Custo de Transbordo ($/t)": 3.0,
                "Permite_Expansao": "SIM",
                "Cap_Expansao_Maxima_Modelo (t)": 75.0,
                "Custo_Expansao_Modelo ($/t)": 900.0,
                "Permite_Granelizacao": "SIM",
                "Custo_Granelizacao_Modelo ($/t)": 520.0,
            },
            {
                "CDA": "W2",
                "Status": "Candidato",
                "Município": "Santos",
                "UF": "SP",
                "Latitude": -23.9535,
                "Longitude": -46.3350,
                "Armazenador": "PRIVADO",
                "Tipo": "Graneleiro",
                "Cap. Estática (t)": 0.0,
                "Cap. Recepção (t)": 0.0,
                "Cap. Expedição (t)": 0.0,
                "Cap. Estática Máxima (t)": 300.0,
                "Custo de Abertura ($)": 3000.0,
                "Custo de Transbordo ($/t)": 7.0,
                "Custo_Fixo_Abertura_Modelo ($)": 1000.0,
                "Custo_Variavel_Capacidade_Modelo ($/t)": 20.0,
                "Permite_Expansao": "NAO",
                "Cap_Expansao_Maxima_Modelo (t)": 0.0,
                "Custo_Expansao_Modelo ($/t)": 0.0,
                "Permite_Granelizacao": "NAO",
                "Custo_Granelizacao_Modelo ($/t)": 0.0,
            },
        ]
    )

    frete = pd.DataFrame(
        [
            {"Estado": "GO", "Frete Tonelada Km": 1.0},
            {"Estado": "SP", "Frete Tonelada Km": 2.0},
        ]
    )

    tarifa_armz = pd.DataFrame(
        [
            {
                "Produto": "Outros",
                "Armazenar_Publico": 50.0,
                "Armazenar_Privado": 60.0,
            },
            {
                "Produto": "Soja",
                "Armazenar_Publico": 45.0,
                "Armazenar_Privado": 55.0,
            },
        ]
    )

    custo_invest = pd.DataFrame(
        [
            {
                "Tipo": "Novo Armazém Convencional",
                "Custo Baixo (R$/t)": 1000.0,
                "Custo Alto (R$/t)": 1500.0,
            },
            {
                "Tipo": "Expansão",
                "Custo Baixo (R$/t)": 600.0,
                "Custo Alto (R$/t)": 1200.0,
            },
            {
                "Tipo": "Granelização",
                "Custo Baixo (R$/t)": 380.0,
                "Custo Alto (R$/t)": 660.0,
            },
        ]
    )

    parametros_modelo = pd.DataFrame(
        [
            {
                "Parametro": "unmet_demand_penalty_default",
                "Valor": 1_000_000.0,
                "Fonte_Parametro": "Sintético",
            }
        ]
    )

    cenarios = pd.DataFrame(
        [
            {
                "Cenario": "base",
                "Probabilidade": 1.0,
                "Ativo": "SIM",
                "Multiplicador_Oferta": 1.0,
                "Multiplicador_Demanda_Domestica": 1.0,
            },
            {
                "Cenario": "baixo",
                "Probabilidade": 0.25,
                "Ativo": "NAO",
                "Multiplicador_Oferta": 0.8,
                "Multiplicador_Demanda_Domestica": 0.9,
            },
            {
                "Cenario": "alto",
                "Probabilidade": 0.25,
                "Ativo": "NAO",
                "Multiplicador_Oferta": 1.2,
                "Multiplicador_Demanda_Domestica": 1.1,
            },
        ]
    )
    oferta_cenarios = pd.DataFrame(
        columns=[
            "Cenario",
            "Produto",
            "Cidade",
            "Data",
            "Peso (ton)",
            "Multiplicador",
        ]
    )
    demanda_cenarios = pd.DataFrame(
        columns=[
            "Cenario",
            "Produto",
            "Cidade",
            "Data",
            "Tipo_Demanda",
            "Peso_Modelo (ton)",
            "Multiplicador",
        ]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        oferta.to_excel(writer, sheet_name="Oferta", index=False)
        demanda.to_excel(writer, sheet_name="Demanda", index=False)
        warehouses.to_excel(writer, sheet_name="Warehouses", index=False)
        frete.to_excel(writer, sheet_name="Frete", index=False)
        tarifa_armz.to_excel(writer, sheet_name="Tarifa_Armz", index=False)
        custo_invest.to_excel(writer, sheet_name="Custo_Invest", index=False)
        parametros_modelo.to_excel(writer, sheet_name="Parametros_Modelo", index=False)
        cenarios.to_excel(writer, sheet_name="Cenarios", index=False)
        oferta_cenarios.to_excel(writer, sheet_name="Oferta_Cenarios", index=False)
        demanda_cenarios.to_excel(writer, sheet_name="Demanda_Cenarios", index=False)


def test_load_model_data_from_golden_excel_schema(tmp_path):
    path = tmp_path / "tiny_golden_agrologistic.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(path)

    assert data.origins == ["Rio Verde - GO", "Jataí - GO"]
    assert data.domestic_customers == ["Goiânia - GO"]
    assert data.export_customers == ["Santos - SP"]

    assert data.warehouses == ["W1", "W2"]
    assert data.existing_warehouses == ["W1"]
    assert data.candidate_warehouses == ["W2"]

    assert data.products == ["Soja"]
    assert data.periods == ["2026-01", "2026-02"]

    assert data.supply[("Rio Verde - GO", "Soja", "2026-01")] == 100.0
    assert data.supply[("Jataí - GO", "Soja", "2026-01")] == 50.0
    assert data.demand_dom[("Goiânia - GO", "Soja", "2026-01")] == 80.0

    assert data.demand_exp[("Santos - SP", "Soja", "2026-01")] == 150.0
    assert data.demand_exp[("Santos - SP", "Soja", "2026-02")] == 120.0

    assert ("Rio Verde - GO", "W1", "Soja") in data.routes_od
    assert ("W1", "Goiânia - GO", "Soja") in data.routes_dc
    assert ("W1", "Santos - SP", "Soja") in data.routes_dc

    assert ("Rio Verde - GO", "W1") in data.dist_od
    assert ("W1", "Goiânia - GO") in data.dist_dc

    assert data.opening_fixed_cost["W2"] == 1000.0
    assert data.candidate_capacity_cost["W2"] == 20.0
    assert data.max_expand_capacity == {"W1": 75.0}
    assert data.expand_fixed_cost == {"W1": 0.0}
    assert data.expand_variable_cost == {"W1": 900.0}
    assert data.max_bulk_capacity == {"W1": 200.0}
    assert data.bulk_fixed_cost == {"W1": 0.0}
    assert data.bulk_variable_cost == {"W1": 520.0}

    assert data.freight_warehouse == {"W1": 1.0, "W2": 2.0}
    assert data.transshipment_cost == {"W1": 3.0, "W2": 7.0}

    assert data.metadata["loader_warnings"] == []
    assert data.scenarios == []


def test_loader_builds_transshipment_routes_distances_and_costs(tmp_path):
    path = tmp_path / "tiny_golden_transshipment.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(include_transshipment_routes=True),
    )

    assert ("W1", "W2", "Soja") in data.routes_dd
    assert ("W2", "W1", "Soja") in data.routes_dd
    assert ("W1", "W1", "Soja") not in data.routes_dd
    assert ("W1", "W2") in data.dist_dd
    assert ("W2", "W1") in data.dist_dd

    assert data.freight_warehouse == {"W1": 1.0, "W2": 2.0}
    assert data.transshipment_cost == {"W1": 3.0, "W2": 7.0}

    result = validate_model_data(
        data=data,
        config=ModelConfig(mode="det", candidate_capacity_mode="scalable"),
        require_distances=True,
    )

    assert result.is_valid, [issue.message for issue in result.errors]


def test_loader_uses_explicit_default_when_transshipment_column_is_absent(tmp_path):
    path = tmp_path / "tiny_golden_transshipment_default.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    all_sheets["Warehouses"] = all_sheets["Warehouses"].drop(
        columns=["Custo de Transbordo ($/t)"]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(default_transshipment_cost=4.5),
    )

    assert data.transshipment_cost == {"W1": 4.5, "W2": 4.5}


def test_loaded_golden_excel_data_passes_model_validation(tmp_path):
    path = tmp_path / "tiny_golden_agrologistic.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(path)

    result = validate_model_data(
        data=data,
        config=ModelConfig(mode="det", candidate_capacity_mode="scalable"),
        require_distances=True,
    )

    assert result.is_valid
    assert result.errors == []


def test_legacy_infinity_demand_is_still_supported(tmp_path):
    path = tmp_path / "legacy_infinity.xlsx"
    build_tiny_golden_excel(path)

    demanda = pd.read_excel(path, sheet_name="Demanda", engine="openpyxl")
    demanda = demanda.drop(
        columns=[
            "Tipo_Demanda",
            "Regra_Limite",
            "Peso_Modelo (ton)",
            "Fonte_Parametro",
            "Observacao",
        ]
    )
    demanda["Peso (ton)"] = demanda["Peso (ton)"].astype(object)
    demanda.loc[1, "Peso (ton)"] = "∞"
    demanda.loc[2, "Peso (ton)"] = "∞"

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    all_sheets["Demanda"] = demanda

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(path)

    assert data.demand_exp[("Santos - SP", "Soja", "2026-01")] == 150.0
    assert data.demand_exp[("Santos - SP", "Soja", "2026-02")] == 120.0
    assert data.metadata["loader_warnings"] != []


def test_fixed_total_candidate_cost_policy_is_used_when_enhanced_columns_absent(tmp_path):
    path = tmp_path / "tiny_golden_agrologistic.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    warehouses = all_sheets["Warehouses"].drop(
        columns=[
            "Custo_Fixo_Abertura_Modelo ($)",
            "Custo_Variavel_Capacidade_Modelo ($/t)",
        ]
    )
    all_sheets["Warehouses"] = warehouses

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(candidate_cost_policy="fixed_total"),
    )

    assert data.opening_fixed_cost["W2"] == 3000.0
    assert data.candidate_capacity_cost["W2"] == 0.0


def test_zero_enhancement_columns_fall_back_to_reported_total_cost(tmp_path):
    path = tmp_path / "zero_candidate_cost_components.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    warehouses = all_sheets["Warehouses"]
    candidate = warehouses["Status"].str.casefold().str.startswith("candidat")
    warehouses.loc[candidate, "Custo_Fixo_Abertura_Modelo ($)"] = 0.0
    warehouses.loc[candidate, "Custo_Variavel_Capacidade_Modelo ($/t)"] = 0.0

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, dataframe in all_sheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(path)

    assert data.opening_fixed_cost["W2"] == 0.0
    assert data.candidate_capacity_cost["W2"] == pytest.approx(10.0)


def test_uncached_candidate_cost_formulas_use_investment_table(tmp_path):
    path = tmp_path / "uncached_candidate_cost_formulas.xlsx"
    build_tiny_golden_excel(path)

    workbook = load_workbook(path)
    warehouses = workbook["Warehouses"]
    columns = {cell.value: cell.column for cell in warehouses[1]}
    max_capacity_column = columns["Cap. Estática Máxima (t)"]
    total_cost_column = columns["Custo de Abertura ($)"]
    fixed_cost_column = columns["Custo_Fixo_Abertura_Modelo ($)"]
    variable_cost_column = columns["Custo_Variavel_Capacidade_Modelo ($/t)"]
    max_capacity_letter = get_column_letter(max_capacity_column)
    total_cost_letter = get_column_letter(total_cost_column)

    warehouses.cell(1, variable_cost_column).value = (
        "Custo_Capacidade_Candidata_Modelo ($/t)"
    )
    warehouses.cell(3, total_cost_column).value = (
        f"={max_capacity_letter}3*AVERAGE(Custo_Invest!$B$2:$C$2)"
    )
    warehouses.cell(3, fixed_cost_column).value = 0.0
    warehouses.cell(3, variable_cost_column).value = (
        f'=IF($B3="Candidato",IF(${max_capacity_letter}3>0,'
        f'${total_cost_letter}3/${max_capacity_letter}3,0),0)'
    )
    workbook.save(path)

    data = load_model_data_from_excel(path)

    assert data.opening_fixed_cost["W2"] == 0.0
    assert data.candidate_capacity_cost["W2"] == pytest.approx(1250.0)
    assert data.metadata["loader_warnings"] == [
        "Derived candidate capacity cost from Custo_Invest for 1 candidate "
        "warehouses because formula-backed cost cells had no cached numeric "
        "values."
    ]


def test_uncached_expansion_formulas_use_parameter_and_investment_tables(tmp_path):
    path = tmp_path / "uncached_expansion_formulas.xlsx"
    build_tiny_golden_excel(path)

    workbook = load_workbook(path)
    parameters = workbook["Parametros_Modelo"]
    parameters.append(
        ["expansion_capacity_fraction_default", 0.25, "Sintético"]
    )

    warehouses = workbook["Warehouses"]
    columns = {cell.value: cell.column for cell in warehouses[1]}
    max_expansion_column = columns["Cap_Expansao_Maxima_Modelo (t)"]
    expansion_cost_column = columns["Custo_Expansao_Modelo ($/t)"]
    warehouses.cell(2, max_expansion_column).value = (
        "=I2*Parametros_Modelo!$B$3"
    )
    warehouses.cell(2, expansion_cost_column).value = (
        "=AVERAGE(Custo_Invest!$B$3:$C$3)"
    )
    workbook.save(path)

    data = load_model_data_from_excel(path)

    assert data.max_expand_capacity["W1"] == pytest.approx(50.0)
    assert data.expand_variable_cost["W1"] == pytest.approx(900.0)
    assert data.metadata["loader_warnings"] == [
        "Derived maximum expansion capacity from Parametros_Modelo for 1 "
        "existing warehouses because formula-backed capacity cells had no "
        "cached numeric values."
    ]


def test_loader_splits_overlapping_domestic_and_export_customers(tmp_path):
    path = tmp_path / "overlap_demand.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    demanda = all_sheets["Demanda"]

    extra_domestic = demanda.iloc[[0]].copy()
    extra_domestic.loc[:, "Cidade"] = "Santos - SP"
    extra_domestic.loc[:, "Latitude"] = -23.9535
    extra_domestic.loc[:, "Longitude"] = -46.3350
    extra_domestic.loc[:, "Tipo_Demanda"] = "DOMESTICA"
    extra_domestic.loc[:, "Regra_Limite"] = "FIXO"
    extra_domestic.loc[:, "Peso (ton)"] = 10.0
    extra_domestic.loc[:, "Peso_Modelo (ton)"] = 10.0
    extra_domestic.loc[:, "Fonte_Parametro"] = "Informado"

    demanda = pd.concat([demanda, extra_domestic], ignore_index=True)
    all_sheets["Demanda"] = demanda

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(path)

    assert "Santos - SP | DOMESTICA" in data.domestic_customers
    assert "Santos - SP | EXPORTACAO" in data.export_customers

    assert set(data.domestic_customers).isdisjoint(set(data.export_customers))

    assert data.demand_dom[("Santos - SP | DOMESTICA", "Soja", "2026-01")] == 10.0
    assert data.demand_exp[("Santos - SP | EXPORTACAO", "Soja", "2026-01")] == 150.0

def test_legacy_infinity_overlap_is_split_with_correct_export_id(tmp_path):
    path = tmp_path / "legacy_overlap_infinity.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    demanda = all_sheets["Demanda"]

    extra_domestic = demanda.iloc[[0]].copy()
    extra_domestic.loc[:, "Cidade"] = "Santos - SP"
    extra_domestic.loc[:, "Latitude"] = -23.9535
    extra_domestic.loc[:, "Longitude"] = -46.3350
    extra_domestic.loc[:, "Tipo_Demanda"] = "DOMESTICA"
    extra_domestic.loc[:, "Regra_Limite"] = "FIXO"
    extra_domestic.loc[:, "Peso (ton)"] = 10.0
    extra_domestic.loc[:, "Peso_Modelo (ton)"] = 10.0

    demanda = pd.concat([demanda, extra_domestic], ignore_index=True)

    demanda = demanda.drop(
        columns=[
            "Tipo_Demanda",
            "Regra_Limite",
            "Peso_Modelo (ton)",
            "Fonte_Parametro",
            "Observacao",
        ]
    )

    demanda["Peso (ton)"] = demanda["Peso (ton)"].astype(object)
    demanda.loc[1, "Peso (ton)"] = "∞"
    demanda.loc[2, "Peso (ton)"] = "∞"

    all_sheets["Demanda"] = demanda

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(path)

    assert "Santos - SP | DOMESTICA" in data.domestic_customers
    assert "Santos - SP | EXPORTACAO" in data.export_customers

    assert set(data.domestic_customers).isdisjoint(set(data.export_customers))

    assert data.demand_dom[("Santos - SP | DOMESTICA", "Soja", "2026-01")] == 10.0
    assert data.demand_exp[("Santos - SP | EXPORTACAO", "Soja", "2026-01")] == 150.0


def test_loader_builds_degenerate_base_scenario_when_enabled(tmp_path):
    path = tmp_path / "base_scenario.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(include_stochastic_scenarios=True),
    )

    assert data.scenarios == ["base"]
    assert data.scenario_prob == {"base": 1.0}
    assert data.supply_s[("base", "Rio Verde - GO", "Soja", "2026-01")] == 100.0
    assert data.demand_dom_s[("base", "Goiânia - GO", "Soja", "2026-01")] == 80.0
    assert data.demand_exp_s[("base", "Santos - SP", "Soja", "2026-01")] == 150.0
    assert data.metadata["scenario_multipliers"] == {
        "base": {
            "supply": 1.0,
            "domestic_demand": 1.0,
            "export_demand": 1.0,
        }
    }


def test_loader_applies_active_scenario_multipliers(tmp_path):
    path = tmp_path / "multiple_scenarios.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    scenarios = all_sheets["Cenarios"]
    scenarios.loc[scenarios["Cenario"] == "base", "Probabilidade"] = 0.5
    scenarios.loc[:, "Ativo"] = "SIM"
    all_sheets["Cenarios"] = scenarios

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(include_stochastic_scenarios=True),
    )

    assert data.scenarios == ["base", "baixo", "alto"]
    assert data.scenario_prob == {"base": 0.5, "baixo": 0.25, "alto": 0.25}
    assert data.supply_s[("baixo", "Rio Verde - GO", "Soja", "2026-01")] == 80.0
    assert data.demand_dom_s[("baixo", "Goiânia - GO", "Soja", "2026-01")] == 72.0
    assert data.demand_exp_s[("baixo", "Santos - SP", "Soja", "2026-01")] == 120.0

    validation = validate_model_data(
        data,
        config=ModelConfig(mode="sto", candidate_capacity_mode="scalable"),
    )
    assert validation.is_valid, [issue.message for issue in validation.errors]


def test_loader_uses_frozen_long_form_workbook_distances(tmp_path):
    path = tmp_path / "frozen_distances.xlsx"
    build_tiny_golden_excel(path)
    sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    rows = []
    pairs = {
        "OD": [
            (origin, warehouse)
            for origin in ("Rio Verde - GO", "Jataí - GO")
            for warehouse in ("W1", "W2")
        ],
        "DC": [
            (warehouse, customer)
            for warehouse in ("W1", "W2")
            for customer in ("Goiânia - GO", "Santos - SP")
        ],
        "DD": [("W1", "W2"), ("W2", "W1")],
    }
    for arc_type, arc_pairs in pairs.items():
        for index, (origin, destination) in enumerate(arc_pairs, start=1):
            rows.append(
                {
                    "Tipo_Arco": arc_type,
                    "Origem": origin,
                    "Destino": destination,
                    "Distancia_km": 1000 + index,
                }
            )
    sheets["Distancias"] = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(
        path,
        ExcelLoaderConfig(
            compute_haversine_distances=False,
            use_workbook_distances=True,
            include_transshipment_routes=True,
        ),
    )

    assert data.dist_od[("Rio Verde - GO", "W1")] == 1001
    assert data.dist_dd[("W2", "W1")] == 1002
    assert data.dist_oc == {}
    assert data.metadata["distance_source"] == "workbook"


def test_loader_rejects_missing_frozen_distance_pairs(tmp_path):
    path = tmp_path / "missing_frozen_distances.xlsx"
    build_tiny_golden_excel(path)

    with pytest.raises(ValueError, match="requires sheet 'Distancias'"):
        load_model_data_from_excel(
            path,
            ExcelLoaderConfig(
                compute_haversine_distances=False,
                use_workbook_distances=True,
            ),
        )


def test_loader_rejects_ambiguous_distance_sources():
    with pytest.raises(ValueError, match="mutually exclusive"):
        ExcelLoaderConfig(use_workbook_distances=True)


def test_loader_rejects_scenario_override_rows_until_supported(tmp_path):
    path = tmp_path / "scenario_overrides.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    all_sheets["Oferta_Cenarios"] = pd.DataFrame(
        [
            {
                "Cenario": "base",
                "Produto": "Soja",
                "Cidade": "Rio Verde - GO",
                "Data": "2026-01",
                "Peso (ton)": 90.0,
                "Multiplicador": 1.0,
            }
        ]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    with pytest.raises(ValueError, match="scenario-specific rows"):
        load_model_data_from_excel(
            path,
            config=ExcelLoaderConfig(include_stochastic_scenarios=True),
        )


def test_loader_ignores_documented_scenario_override_placeholders(tmp_path):
    path = tmp_path / "scenario_override_placeholders.xlsx"
    build_tiny_golden_excel(path)

    all_sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    all_sheets["Oferta_Cenarios"] = pd.DataFrame(
        [
            {
                "Cenario": "base",
                "Produto": None,
                "Cidade": None,
                "Data": None,
                "Peso (ton)": None,
                "Multiplicador": 1.0,
                "Fonte_Parametro": "Informado",
                "Observacao": "Optional template row.",
            }
        ]
    )
    all_sheets["Demanda_Cenarios"] = pd.DataFrame(
        [
            {
                "Cenario": "base",
                "Produto": None,
                "Cidade": None,
                "Data": None,
                "Tipo_Demanda": "DOMESTICA",
                "Regra_Limite": "FIXO",
                "Peso (ton)": None,
                "Peso_Modelo (ton)": None,
                "Multiplicador": 1.0,
                "Fonte_Parametro": "Informado",
                "Observacao": "Optional template row.",
            }
        ]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, dataframe in all_sheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(include_stochastic_scenarios=True),
    )

    assert data.scenarios == ["base"]


def test_loader_generates_all_nine_supply_demand_combinations(tmp_path):
    path = tmp_path / "nine_scenarios.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(
            include_stochastic_scenarios=True,
            scenario_generation_mode="cartesian",
        ),
    )

    assert len(data.scenarios) == 9
    assert sum(data.scenario_prob.values()) == pytest.approx(1.0)
    scenario = "oferta_baixo__demanda_alto"
    assert scenario in data.scenarios
    assert data.supply_s[(scenario, "Rio Verde - GO", "Soja", "2026-01")] == 80.0
    assert data.demand_dom_s[(scenario, "Goiânia - GO", "Soja", "2026-01")] == 88.0
    assert data.demand_exp_s[(scenario, "Santos - SP", "Soja", "2026-01")] == 120.0


def test_loader_accepts_an_explicit_user_selected_scenario_subset(tmp_path):
    path = tmp_path / "selected_scenarios.xlsx"
    build_tiny_golden_excel(path)

    data = load_model_data_from_excel(
        path,
        config=ExcelLoaderConfig(
            include_stochastic_scenarios=True,
            scenario_generation_mode="cartesian",
            stochastic_combinations=(
                ("baixo", "alto"),
                ("base", "base"),
                ("alto", "baixo"),
            ),
            stochastic_probabilities=(0.2, 0.5, 0.3),
        ),
    )

    assert data.scenarios == [
        "oferta_baixo__demanda_alto",
        "oferta_base__demanda_base",
        "oferta_alto__demanda_baixo",
    ]
    assert data.scenario_prob == {
        "oferta_baixo__demanda_alto": 0.2,
        "oferta_base__demanda_base": 0.5,
        "oferta_alto__demanda_baixo": 0.3,
    }

    validation = validate_model_data(
        data,
        config=ModelConfig(mode="sto", candidate_capacity_mode="scalable"),
    )
    assert validation.is_valid, [issue.message for issue in validation.errors]

